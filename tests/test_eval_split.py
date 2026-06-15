import importlib.util
from pathlib import Path

import pytest
from docx import Document as DocxDocument

from genacademy_coach.eval_split import (
    assert_normalized_eval_filename,
    ngrams,
    phrase_hashes,
    split_items,
)


def test_rejects_case_variant_eval_filename():
    with pytest.raises(ValueError, match="lowercase kebab-case"):
        assert_normalized_eval_filename(Path("Week1_Session2_Chat_Questions.docx"))


def test_split_items_is_deterministic_and_keeps_test_frozen():
    items = [{"id": f"q{i}", "source_sha256": f"s{i}"} for i in range(100)]

    first = split_items(items, seed="genacademy-coach-v1")
    second = split_items(items, seed="genacademy-coach-v1")
    expanded = split_items(
        [*items, {"id": "q-new", "source_sha256": "s-new"}],
        seed="genacademy-coach-v1",
    )
    first_by_id = {item["id"]: item["split"] for item in first}

    assert first == second
    assert {row["split"] for row in first} == {"seed", "dev", "test"}
    assert all(
        row["split"] == first_by_id[row["id"]]
        for row in expanded
        if row["id"] in first_by_id
    )
    assert all("text" not in row for row in first)


def test_ngrams_uses_all_normalized_eight_word_phrases_without_private_manifest_text():
    text = "One two three four five six seven eight nine ten"

    phrases = ngrams(text, n=8)

    assert "one two three four five six seven eight" in phrases
    assert "two three four five six seven eight nine" in phrases
    assert "three four five six seven eight nine ten" in phrases


def test_phrase_hashes_preserve_multiple_eval_sources():
    phrases = phrase_hashes(
        [
            ("first.md", "One two three four five six seven eight"),
            ("second.md", "One two three four five six seven eight"),
        ],
        n=8,
    )

    assert len(phrases["one two three four five six seven eight"]) == 2
    assert {row["source_file"] for row in phrases["one two three four five six seven eight"]} == {
        "first.md",
        "second.md",
    }


def test_read_eval_text_includes_docx_tables(tmp_path):
    script_path = Path("scripts/check_eval_leak.py").resolve()
    spec = importlib.util.spec_from_file_location("check_eval_leak", script_path)
    assert spec is not None
    assert spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    path = tmp_path / "question-table.docx"
    doc = DocxDocument()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Learner question"
    table.cell(0, 1).text = "Expected concept"
    doc.save(path)

    text = module.read_eval_text(path)

    assert "Learner question" in text
    assert "Expected concept" in text
